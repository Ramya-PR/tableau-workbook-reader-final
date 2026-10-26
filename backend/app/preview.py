"""Row-level data preview for data files bundled in a .twbx.

* ``.hyper``  -> read with Tableau's Hyper API (optional dependency).
* ``.csv`` / ``.tsv`` / ``.txt`` / ``.tab`` -> read with the csv module.
* ``.xlsx`` / ``.xlsm`` -> read with openpyxl in read-only (streaming) mode; one table per sheet.
* anything else (``.tde``, legacy ``.xls``, ...) is listed but not previewable.

User-supplied table names are never interpolated into SQL: a requested table
must match one of the names returned by the file's own catalog, and the SQL is
built from the catalog's ``TableName`` object.
"""

from __future__ import annotations

import csv
import datetime as dt
import decimal
import itertools
import math
import zipfile
from pathlib import Path
from typing import Any

from .archive import validate_zip
from .config import Settings
from .errors import AppError, not_found, parse_error

HYPER_INSTALL_HINT = (
    "Previewing .hyper extracts needs Tableau's Hyper API. Install it with: pip install tableauhyperapi"
)


def hyper_available() -> bool:
    try:
        import tableauhyperapi  # noqa: F401
    except Exception:  # ImportError, or a platform the wheel doesn't support
        return False
    return True


XLSX_SUFFIXES = (".xlsx", ".xlsm")
MAX_COUNTED_ROWS = 1_000_000  # stop counting rows in a sheet here (keeps huge sheets from stalling the request)


def describe_file(kind: str, name: str = "") -> tuple[bool, str | None]:
    """Is a data file of this kind previewable, and if not, why not?"""
    if kind == "hyper":
        return (True, None) if hyper_available() else (False, HYPER_INSTALL_HINT)
    if kind == "csv":
        return True, None
    if kind == "tde":
        return False, (
            "Legacy .tde extract. Open the workbook in Tableau and re-save it (Tableau converts extracts to .hyper)."
        )
    if kind == "excel":
        if name and not name.lower().endswith(XLSX_SUFFIXES):
            return False, "Legacy .xls files are not previewed. Re-save the file as .xlsx to preview it."
        return True, None
    return False, "This file type is not previewed in this version."


# --------------------------------------------------------------------------
# JSON-safe values
# --------------------------------------------------------------------------
def json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, str)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, decimal.Decimal):
        return str(value)  # keep full precision
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, (bytes, bytearray)):
        return f"<{len(value)} bytes>"
    # Hyper API value types (Timestamp, Date, Interval, ...)
    for attr in ("to_datetime", "to_date"):
        conv = getattr(value, attr, None)
        if callable(conv):
            try:
                return conv().isoformat()
            except Exception:  # e.g. year out of Python's range
                break
    return str(value)


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------
def _open_csv(path: Path):
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            with open(path, encoding=encoding, newline="") as fh:
                fh.read(1 << 20)  # decode-check the first MB
            return open(path, encoding=encoding, newline="")
        except UnicodeDecodeError:
            continue
    return open(path, encoding="latin-1", newline="")


def _csv_dialect(path: Path):
    with _open_csv(path) as fh:
        sample = fh.read(8192)
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        return csv.excel


def _csv_tables(path: Path) -> list[dict]:
    dialect = _csv_dialect(path)
    with _open_csv(path) as fh:
        reader = csv.reader(fh, dialect)
        header = next(reader, None)
        if header is None:
            return [{"name": path.name, "row_count": 0, "columns": []}]
        count = sum(1 for _ in reader)
    return [
        {
            "name": path.name,
            "row_count": count,
            "columns": [{"name": h or f"column_{i + 1}", "type": "text"} for i, h in enumerate(header)],
        }
    ]


def _csv_rows(path: Path, offset: int, limit: int) -> dict:
    dialect = _csv_dialect(path)
    with _open_csv(path) as fh:
        reader = csv.reader(fh, dialect)
        header = next(reader, None) or []
        width = len(header)
        rows = []
        for raw in itertools.islice(reader, offset, offset + limit):
            row = (raw + [None] * width)[:width] if width else raw
            rows.append(row)
    return {
        "columns": [{"name": h or f"column_{i + 1}", "type": "text"} for i, h in enumerate(header)],
        "rows": rows,
    }


# --------------------------------------------------------------------------
# Excel (.xlsx / .xlsm)
# --------------------------------------------------------------------------
def _open_workbook(path: Path):
    """Open an .xlsx for streaming reads, after checking the inner zip with the same guards as a .twbx."""
    if not path.name.lower().endswith(XLSX_SUFFIXES):
        raise AppError("PREVIEW_UNAVAILABLE", describe_file("excel", path.name)[1] or "Not previewable.", 501)
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(path) as zf:
            validate_zip(zf, Settings())
        return load_workbook(path, read_only=True, data_only=True)
    except AppError:
        raise
    except Exception as exc:  # BadZipFile, KeyError, InvalidFileException, ...
        raise parse_error("The Excel file could not be read.", "EXCEL_READ_ERROR") from exc


def _header(row) -> list[dict]:
    return [{"name": str(v) if v not in (None, "") else f"column_{i + 1}", "type": "text"} for i, v in enumerate(row)]


def _excel_tables(path: Path) -> list[dict]:
    wb = _open_workbook(path)
    try:
        out = []
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            first = next(rows, None)
            count = sum(1 for _ in itertools.islice(rows, MAX_COUNTED_ROWS)) if first is not None else 0
            out.append({"name": ws.title, "row_count": count, "columns": _header(first or [])})
        return out
    finally:
        wb.close()


def _excel_rows(path: Path, table: str | None, offset: int, limit: int) -> dict:
    wb = _open_workbook(path)
    try:
        if table not in wb.sheetnames:
            raise not_found(f"Sheet {table!r} was not found in this workbook.", "UNKNOWN_TABLE")
        rows = wb[table].iter_rows(values_only=True)
        first = next(rows, None) or ()
        width = len(first)
        page = [
            [json_safe(v) for v in (list(r) + [None] * width)[:width]]
            for r in itertools.islice(rows, offset, offset + limit)
        ]
        return {"columns": _header(first), "rows": page}
    finally:
        wb.close()


# --------------------------------------------------------------------------
# Hyper
# --------------------------------------------------------------------------
def _hyper_table_label(table_name) -> str:
    schema = getattr(table_name, "schema_name", None)
    name = table_name.name.unescaped
    return f"{schema.name.unescaped}.{name}" if schema is not None else name


def _with_hyper(path: Path, fn):
    if not hyper_available():
        raise AppError("PREVIEW_UNAVAILABLE", HYPER_INSTALL_HINT, 501)
    from tableauhyperapi import Connection, HyperException, HyperProcess, Telemetry

    try:
        # log_config="" stops hyperd from writing a log file into the cwd.
        with HyperProcess(
            telemetry=Telemetry.DO_NOT_SEND_USAGE_DATA_TO_TABLEAU, parameters={"log_config": ""}
        ) as hyper:
            with Connection(hyper.endpoint, database=str(path)) as conn:
                return fn(conn)
    except HyperException as exc:
        raise parse_error(f"The .hyper file could not be read: {exc.main_message}", "HYPER_READ_ERROR") from exc


def _hyper_catalog(conn) -> list:
    tables = []
    for schema in conn.catalog.get_schema_names():
        tables.extend(conn.catalog.get_table_names(schema))
    return tables


def _hyper_tables(path: Path) -> list[dict]:
    def run(conn):
        out = []
        for tn in _hyper_catalog(conn):
            definition = conn.catalog.get_table_definition(tn)
            out.append(
                {
                    "name": _hyper_table_label(tn),
                    "row_count": conn.execute_scalar_query(f"SELECT COUNT(*) FROM {tn}"),
                    "columns": [{"name": c.name.unescaped, "type": str(c.type)} for c in definition.columns],
                }
            )
        return out

    return _with_hyper(path, run)


def _hyper_rows(path: Path, table: str, offset: int, limit: int) -> dict:
    def run(conn):
        match = next((tn for tn in _hyper_catalog(conn) if _hyper_table_label(tn) == table), None)
        if match is None:
            raise not_found(f"Table {table!r} was not found in this extract.", "UNKNOWN_TABLE")
        definition = conn.catalog.get_table_definition(match)
        # offset/limit are validated ints; the table comes from the catalog, not from user text
        result = conn.execute_list_query(f"SELECT * FROM {match} LIMIT {int(limit)} OFFSET {int(offset)}")
        return {
            "columns": [{"name": c.name.unescaped, "type": str(c.type)} for c in definition.columns],
            "rows": [[json_safe(v) for v in row] for row in result],
        }

    return _with_hyper(path, run)


# --------------------------------------------------------------------------
# public API
# --------------------------------------------------------------------------
def list_tables(path: Path, kind: str) -> list[dict]:
    if kind == "hyper":
        return _hyper_tables(path)
    if kind == "csv":
        return _csv_tables(path)
    if kind == "excel":
        return _excel_tables(path)
    raise AppError("PREVIEW_UNAVAILABLE", describe_file(kind)[1] or "Not previewable.", 501)


def read_rows(path: Path, kind: str, table: str | None, offset: int, limit: int) -> dict:
    if kind == "hyper":
        if not table:
            raise parse_error("A table name is required to preview a .hyper extract.", "TABLE_REQUIRED")
        data = _hyper_rows(path, table, offset, limit)
    elif kind == "csv":
        data = _csv_rows(path, offset, limit)
    elif kind == "excel":
        data = _excel_rows(path, table, offset, limit)
    else:
        raise AppError("PREVIEW_UNAVAILABLE", describe_file(kind)[1] or "Not previewable.", 501)
    data.update({"offset": offset, "limit": limit})
    return data
